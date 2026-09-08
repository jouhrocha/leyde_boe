#!/usr/bin/env python3
"""
word_a_txt.py
Selecciona varios archivos .docx con una ventana gráfica
y vuelca su contenido en texto plano dentro de un único .txt.
"""

import os
import sys
import tkinter as tk
from tkinter import filedialog, messagebox, ttk

# ── dependencias ──────────────────────────────────────────────────────────────
try:
    from docx import Document
except ImportError:
    import subprocess
    print("Instalando python-docx...")
    subprocess.check_call([sys.executable, "-m", "pip", "install", "python-docx"])
    from docx import Document

# ── separador visual ──────────────────────────────────────────────────────────
SEP = "=" * 72


def extraer_texto(ruta_docx: str) -> str:
    """Extrae todo el texto de un .docx (párrafos + tablas)."""
    doc = Document(ruta_docx)
    lineas = []

    # Párrafos normales
    for p in doc.paragraphs:
        texto = p.text.strip()
        if texto:
            lineas.append(texto)

    # Tablas
    for tabla in doc.tables:
        for fila in tabla.rows:
            fila_txt = "  |  ".join(
                celda.text.strip() for celda in fila.cells if celda.text.strip()
            )
            if fila_txt:
                lineas.append(fila_txt)

    return "\n".join(lineas)


def procesar(archivos: list[str], ruta_salida: str, barra, lbl_estado):
    """Procesa cada archivo y escribe el .txt de salida."""
    total = len(archivos)
    bloques = []

    for i, ruta in enumerate(archivos, 1):
        nombre = os.path.basename(ruta)
        lbl_estado.config(text=f"Procesando: {nombre}")
        barra["value"] = (i / total) * 100
        barra.update()

        try:
            contenido = extraer_texto(ruta)
            titulo = os.path.splitext(nombre)[0]
            bloque = (
                f"{SEP}\n"
                f"  DOCUMENTO: {titulo}\n"
                f"{SEP}\n\n"
                f"{contenido}\n"
            )
            bloques.append(bloque)
        except Exception as e:
            bloques.append(
                f"{SEP}\n"
                f"  DOCUMENTO: {nombre}  [ERROR: {e}]\n"
                f"{SEP}\n\n"
            )

    with open(ruta_salida, "w", encoding="utf-8") as f:
        f.write("\n\n".join(bloques))

    lbl_estado.config(text=f"✅ Listo  →  {ruta_salida}")
    barra["value"] = 100
    messagebox.showinfo(
        "Completado",
        f"{total} documento(s) exportados.\n\nArchivo guardado en:\n{ruta_salida}",
    )


# ── GUI principal ─────────────────────────────────────────────────────────────
class App(tk.Tk):
    def __init__(self):
        super().__init__()
        self.title("Word → TXT  |  Extractor de texto plano")
        self.resizable(False, False)
        self.configure(bg="#f5f5f5")
        self._archivos: list[str] = []
        self._construir_ui()

    def _construir_ui(self):
        PAD = dict(padx=16, pady=8)

        # ── título ────────────────────────────────────────────────────────────
        tk.Label(
            self,
            text="Word  →  TXT",
            font=("Segoe UI", 16, "bold"),
            bg="#f5f5f5",
            fg="#2c3e50",
        ).grid(row=0, column=0, columnspan=2, pady=(18, 2))

        tk.Label(
            self,
            text="Selecciona archivos .docx y extrae su contenido a texto plano.",
            font=("Segoe UI", 9),
            bg="#f5f5f5",
            fg="#666",
        ).grid(row=1, column=0, columnspan=2, pady=(0, 12))

        # ── botón seleccionar ─────────────────────────────────────────────────
        tk.Button(
            self,
            text="📂  Seleccionar archivos .docx",
            font=("Segoe UI", 10),
            bg="#2980b9",
            fg="white",
            relief="flat",
            cursor="hand2",
            padx=12,
            pady=6,
            command=self._seleccionar,
        ).grid(row=2, column=0, columnspan=2, **PAD)

        # ── lista de archivos seleccionados ───────────────────────────────────
        frame_lista = tk.Frame(self, bg="#f5f5f5")
        frame_lista.grid(row=3, column=0, columnspan=2, padx=16, pady=4, sticky="ew")

        scrollbar = tk.Scrollbar(frame_lista)
        scrollbar.pack(side="right", fill="y")

        self.listbox = tk.Listbox(
            frame_lista,
            width=60,
            height=8,
            font=("Consolas", 9),
            selectmode="extended",
            yscrollcommand=scrollbar.set,
            bg="white",
            fg="#333",
            relief="flat",
            bd=1,
            highlightthickness=1,
            highlightbackground="#ccc",
        )
        self.listbox.pack(side="left", fill="both")
        scrollbar.config(command=self.listbox.yview)

        # ── botón quitar seleccionados ────────────────────────────────────────
        tk.Button(
            self,
            text="🗑  Quitar seleccionados",
            font=("Segoe UI", 9),
            bg="#e74c3c",
            fg="white",
            relief="flat",
            cursor="hand2",
            padx=10,
            pady=4,
            command=self._quitar,
        ).grid(row=4, column=0, columnspan=2, pady=(2, 8))

        # ── nombre archivo salida ─────────────────────────────────────────────
        tk.Label(
            self, text="Nombre del archivo de salida:", bg="#f5f5f5", font=("Segoe UI", 9)
        ).grid(row=5, column=0, sticky="e", padx=(16, 4))

        self.entry_salida = tk.Entry(self, width=28, font=("Segoe UI", 9))
        self.entry_salida.insert(0, "documentos_combinados.txt")
        self.entry_salida.grid(row=5, column=1, sticky="w", padx=(0, 16), pady=4)

        # ── barra de progreso ─────────────────────────────────────────────────
        self.barra = ttk.Progressbar(self, length=440, mode="determinate")
        self.barra.grid(row=6, column=0, columnspan=2, padx=16, pady=(8, 4))

        self.lbl_estado = tk.Label(
            self, text="Esperando archivos…", bg="#f5f5f5", fg="#555", font=("Segoe UI", 9)
        )
        self.lbl_estado.grid(row=7, column=0, columnspan=2, pady=(0, 6))

        # ── botón ejecutar ────────────────────────────────────────────────────
        tk.Button(
            self,
            text="⚡  Extraer a TXT",
            font=("Segoe UI", 11, "bold"),
            bg="#27ae60",
            fg="white",
            relief="flat",
            cursor="hand2",
            padx=18,
            pady=8,
            command=self._ejecutar,
        ).grid(row=8, column=0, columnspan=2, pady=(6, 20))

    # ── acciones ──────────────────────────────────────────────────────────────
    def _seleccionar(self):
        rutas = filedialog.askopenfilenames(
            title="Selecciona archivos Word",
            filetypes=[("Word Documents", "*.docx"), ("Todos los archivos", "*.*")],
        )
        for r in rutas:
            if r not in self._archivos:
                self._archivos.append(r)
                self.listbox.insert("end", os.path.basename(r))
        self.lbl_estado.config(text=f"{len(self._archivos)} archivo(s) en cola.")

    def _quitar(self):
        seleccionados = list(self.listbox.curselection())[::-1]
        for i in seleccionados:
            self.listbox.delete(i)
            del self._archivos[i]
        self.lbl_estado.config(text=f"{len(self._archivos)} archivo(s) en cola.")

    def _ejecutar(self):
        if not self._archivos:
            messagebox.showwarning("Sin archivos", "Añade al menos un archivo .docx.")
            return

        nombre_base = self.entry_salida.get().strip() or "documentos_combinados.txt"
        if not nombre_base.endswith(".txt"):
            nombre_base += ".txt"

        ruta_salida = filedialog.asksaveasfilename(
            title="Guardar archivo TXT como…",
            defaultextension=".txt",
            initialfile=nombre_base,
            filetypes=[("Archivo de texto", "*.txt")],
        )
        if not ruta_salida:
            return

        self.barra["value"] = 0
        procesar(self._archivos, ruta_salida, self.barra, self.lbl_estado)


# ── punto de entrada ──────────────────────────────────────────────────────────
if __name__ == "__main__":
    app = App()
    app.mainloop()
